"""Excel eksport regressiyasi (2026-10-20): test bazasidagi HAR BIR KP (oxirgi revision) eksport qilinadi va
fayl ochilib, qatorlar va jami summalar bazadagi hisob (v_line_price / v_annex_totals) bilan solishtiriladi.
Narxi noma'lum qatori bor KP eksport qilinmasligi (xato yashirilmasligi) ham tekshiriladi."""
import io
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import text

from backend.app import db as app_db
from backend.app import kp_export
from backend.app.kp_export import BLOCK_SIZE, BLOCK_START, LINES_PER_BLOCK


@pytest.fixture(scope="module")
def ro_conn():
    with app_db.engine().connect() as c:
        c.exec_driver_sql("SET TRANSACTION READ ONLY")       # eksport faqat o'qiydi
        yield c
        c.rollback()


def _dec(v):
    return None if v is None else Decimal(str(v))


def _same(excel, db) -> bool:
    """Excel sonni 64-bit (taxminan 15 xona) aniqlikda saqlaydi; bazada aniqlik to'liq (masalan 37348.000000000004).
    Farq Excel aniqligi chegarasidan (nisbiy 1e-12) kichik bo'lsa — bir xil qiymat."""
    if excel is None or db is None:
        return excel is None and db is None
    return abs(_dec(excel) - Decimal(db)) <= abs(Decimal(db)) * Decimal("1e-12") + Decimal("1e-9")


def test_every_priced_kp_exports_and_totals_match(ro_conn):
    kps = ro_conn.execute(text(
        """SELECT q.kp_number, t.revision_id FROM quotation q
           JOIN LATERAL (SELECT r.id FROM quotation_revision r WHERE r.quotation_id = q.id ORDER BY r.revision_no DESC LIMIT 1) r ON true
           JOIN v_quotation_totals t ON t.revision_id = r.id
           WHERE t.subtotal IS NOT NULL ORDER BY q.kp_number""")).all()
    assert len(kps) > 300
    problems = []
    for kp, rev in kps:
        try:
            _, data = kp_export.build(ro_conn, kp, None)
        except kp_export.ExportError as e:                   # masalan shablon chegarasi — yashirilmaydi, ro'yxatga
            problems.append(f"{kp}: {e}")
            continue
        ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
        hidden = {r for r, d in ws.row_dimensions.items() if d.hidden}
        annexes = ro_conn.execute(text("SELECT annex_id, subtotal FROM v_annex_totals WHERE revision_id = :r ORDER BY annex_no"),
                                  {"r": rev}).all()
        for b, (annex_id, subtotal) in enumerate(annexes):
            s = BLOCK_START + BLOCK_SIZE * b
            lines = ro_conn.execute(text(
                "SELECT model, quantity, display_line_total FROM v_line_price WHERE annex_id = :a ORDER BY line_no"),
                {"a": annex_id}).all()
            for i, ln in enumerate(lines):
                r = s + 4 + i
                if r in hidden or not _same(ws[f"I{r}"].value, ln.display_line_total) or _dec(ws[f"F{r}"].value) != ln.quantity:
                    problems.append(f"{kp} ilova {b + 1} qator {i + 1}: Excel I={ws[f'I{r}'].value} F={ws[f'F{r}'].value}"
                                    f" baza={ln.display_line_total}/{ln.quantity}")
                    break
            if s + 4 + len(lines) < s + 4 + LINES_PER_BLOCK and s + 4 + len(lines) not in hidden:
                problems.append(f"{kp} ilova {b + 1}: bo'sh qator yashirilmagan")
            if not _same(ws[f"H{s + 82}"].value, subtotal):
                problems.append(f"{kp} ilova {b + 1}: Итого Excel={ws[f'H{s + 82}'].value} baza={subtotal}")
    assert not problems, "\n".join(problems[:30]) + f"\n... jami {len(problems)}"


def test_kp_with_unknown_price_is_not_exported(ro_conn):
    """9 ta narxsiz qator (5 KP) — eksport aniq xato bilan to'xtaydi, narx taxmin qilinmaydi."""
    kps = ro_conn.execute(text(
        """SELECT DISTINCT q.kp_number FROM v_line_price v JOIN quotation_annex a ON a.id = v.annex_id
           JOIN quotation_revision r ON r.id = a.revision_id JOIN quotation q ON q.id = r.quotation_id
           WHERE v.display_line_total IS NULL""")).scalars().all()
    assert kps
    for kp in kps:
        with pytest.raises(kp_export.ExportError, match="narxi noma'lum"):
            kp_export.build(ro_conn, kp, None)
